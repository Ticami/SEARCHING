from datetime import datetime, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook

from searching import relatorios
from searching.db.models import Alerta, Coleta, Execucao, Produto, RevisaoLink
from searching.db.session import criar_engine, iniciar_banco

AGORA = datetime(2026, 10, 23, 15, 0)


@pytest.fixture
def s(tmp_path):
    with iniciar_banco(criar_engine(tmp_path / "t.db"))() as sessao:
        yield sessao


def _produto(s, nome, loja="Magazine Luiza", ativo=True):
    p = Produto(nome_comercial=nome, loja=loja, url_atual=f"https://x/{nome.replace(' ', '-')}/p/1/", ativo=ativo)
    s.add(p)
    s.flush()
    return p


def _coleta(s, p, quando, status="OK", pix=None, cartao=None, print_=None, exec_id=None, erro=None):
    c = Coleta(produto_id=p.id, execucao_id=exec_id, coletado_em=quando, status=status, preco_pix_centavos=pix,
               preco_cartao_1x_centavos=cartao, caminho_print=print_, erro_detalhe=erro, titulo_lido=f"Título de {p.nome_comercial}")
    s.add(c)
    s.flush()
    return c


@pytest.fixture
def cenario(s):
    a, b, amazon = _produto(s, "Galaxy S25"), _produto(s, "Galaxy A17"), _produto(s, "Galaxy S24", loja="Amazon")
    c = {
        "a1": _coleta(s, a, datetime(2026, 10, 21, 12, 0), pix=400000, cartao=430000, print_="prints/2026-10-21/a.png"),
        "a2": _coleta(s, a, datetime(2026, 10, 22, 12, 0), pix=380000, cartao=410000, print_="prints/2026-10-22/a.png"),
        "b1": _coleta(s, b, datetime(2026, 10, 22, 12, 1), status="BLOQUEIO", erro="página de verificação"),
        "m1": _coleta(s, amazon, datetime(2026, 10, 22, 12, 2), pix=300000),
    }
    s.commit()
    return {"a": a, "b": b, "amazon": amazon, **c}


# ---------- histórico ----------
def test_historico_traz_texto_formatado_e_valores_numericos(s, cenario):
    linha = relatorios.historico(s, produto_id=cenario["a"].id)[0]
    assert linha["PIX"] == "R$ 4.000,00" and linha["PIX_CENTAVOS"] == 400000 and linha["CREDITO_CENTAVOS"] == 430000
    assert linha["DATA"] == "2026-10-21 12:00:00" and linha["SITE"] == "Magazine Luiza" and linha["STATUS"] == "OK"
    assert linha["ID"] == cenario["a1"].id


def test_filtros_de_periodo_loja_produto_e_status(s, cenario):
    assert len(relatorios.historico(s)) == 4
    assert [r["DATA"][:10] for r in relatorios.historico(s, de="2026-10-22")] == ["2026-10-22"] * 3  # inclui o dia inicial
    assert len(relatorios.historico(s, ate="2026-10-21")) == 1  # inclui o dia final inteiro
    assert len(relatorios.historico(s, de="2026-10-22", ate="2026-10-22")) == 3
    assert len(relatorios.historico(s, loja="Amazon")) == 1
    assert [r["STATUS"] for r in relatorios.historico(s, status="BLOQUEIO")] == ["BLOQUEIO"]
    assert relatorios.historico(s, de="2026-11-01") == []
    with pytest.raises(ValueError):
        relatorios.historico(s, de="23/10/2026")


# ---------- galeria ----------
def test_galeria_por_data_e_por_produto_com_preco_ao_lado(s, cenario):
    assert relatorios.datas_com_prints(s) == [{"data": "2026-10-22", "total": 3, "com_print": 1}, {"data": "2026-10-21", "total": 1, "com_print": 1}]
    itens = relatorios.galeria(s, data="2026-10-22")
    assert len(itens) == 1 and itens[0]["pix"] == "R$ 3.800,00" and itens[0]["cartao_1x"] == "R$ 4.100,00"
    assert itens[0]["print_url"] == "/api/prints/2026-10-22/a.png"
    por_produto = relatorios.galeria(s, produto_id=cenario["a"].id)
    assert [i["pix"] for i in por_produto] == ["R$ 3.800,00", "R$ 4.000,00"]  # mais recente primeiro
    assert relatorios.galeria(s, produto_id=cenario["b"].id) == []  # sem print


def test_url_do_print_e_codificada_e_aceita_caminho_absoluto():
    assert relatorios.print_url("prints/2025-12-08_16-19-01/Samsung Galaxy S24+ 256GB_16-19-05.png") == \
        "/api/prints/2025-12-08_16-19-01/Samsung%20Galaxy%20S24%2B%20256GB_16-19-05.png"
    assert relatorios.print_url("C:\\tmp\\prints\\2026-10-23\\x.png") == "/api/prints/2026-10-23/x.png"
    assert relatorios.print_url(None) is None and relatorios.print_url("sozinho.png") is None


# ---------- Excel ----------
def test_excel_tem_precos_numericos_e_caminho_do_print(s, cenario):
    ids = [cenario["a1"].id, cenario["b1"].id]
    ws = load_workbook(BytesIO(relatorios.exportar_excel(s, ids))).active
    cab = [c.value for c in ws[1]]
    assert cab[:2] == ["Data", "Produto"] and "Print (caminho)" in cab
    linhas = {r[1].value: r for r in ws.iter_rows(min_row=2)}
    a = linhas["Galaxy S25"]
    assert a[6].value == 4000.0 and a[7].value == 4300.0 and isinstance(a[6].value, (int, float))  # números, não texto
    assert a[6].number_format == '"R$" #,##0.00' and a[0].value == datetime(2026, 10, 21, 12, 0)
    assert a[9].value == "prints/2026-10-21/a.png"
    b = linhas["Galaxy A17"]
    assert b[5].value == "BLOQUEIO" and b[6].value is None and b[10].value == "página de verificação"
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref


def test_excel_valida_a_lista(s, cenario):
    with pytest.raises(ValueError):
        relatorios.exportar_excel(s, [])


# ---------- saúde ----------
def test_saude_por_loja_com_uma_coleta_final_por_produto_e_execucao(s):
    amz = _produto(s, "Galaxy S24", loja="Amazon")
    a, b = _produto(s, "Galaxy S25"), _produto(s, "Galaxy A17")
    e1, e2 = Execucao(iniciada_em=AGORA, origem="agendada"), Execucao(iniciada_em=AGORA, origem="agendada")
    s.add_all([e1, e2])
    s.flush()
    _coleta(s, a, datetime(2026, 10, 22, 12), pix=1, exec_id=e1.id)
    _coleta(s, b, datetime(2026, 10, 22, 12, 1), status="LINK_QUEBRADO", exec_id=e1.id)
    _coleta(s, b, datetime(2026, 10, 22, 12, 5), pix=1, exec_id=e1.id)  # refeita depois de corrigir o link: vale esta
    _coleta(s, a, datetime(2026, 10, 23, 12), status="BLOQUEIO", exec_id=e2.id)
    _coleta(s, b, datetime(2026, 10, 23, 12, 1), status="BLOQUEIO", exec_id=e2.id)
    _coleta(s, amz, datetime(2026, 10, 23, 12, 2), pix=1, exec_id=e2.id)
    _coleta(s, a, datetime(2026, 10, 1, 12), status="TIMEOUT", exec_id=e1.id)  # fora dos 7 dias
    s.add(RevisaoLink(produto_id=b.id, motivo="LINK_QUEBRADO", candidatos=[]))
    s.add(Alerta(produto_id=a.id, preco_anterior_centavos=2, preco_atual_centavos=1, queda_pct=50, enviado_em=AGORA, sucesso=True))
    s.add(Alerta(produto_id=a.id, preco_anterior_centavos=2, preco_atual_centavos=1, queda_pct=50, enviado_em=AGORA, sucesso=False))
    s.commit()

    r = relatorios.saude(s, agora=AGORA)
    lojas = {l["loja"]: l for l in r["lojas"]}
    assert (lojas["Magazine Luiza"]["total"], lojas["Magazine Luiza"]["ok"], lojas["Magazine Luiza"]["taxa_sucesso"]) == (4, 2, 50.0)
    assert (lojas["Amazon"]["total"], lojas["Amazon"]["taxa_sucesso"]) == (1, 100.0)
    assert r["falhas_por_status"] == {"BLOQUEIO": 2}  # a falha refeita e a de 22 dias atrás não entram
    assert {p["produto"] for p in r["produtos_com_falha_na_ultima_coleta"]} == {"Galaxy S25", "Galaxy A17"}
    assert r["links_em_revisao"] == 1 and r["alertas"] == {"enviados": 1, "falhos": 1}


def test_saude_sem_dados(s):
    r = relatorios.saude(s, agora=AGORA)
    assert r["lojas"] == [] and r["falhas_por_status"] == {} and r["alertas"] == {"enviados": 0, "falhos": 0}


def test_ultima_execucao_e_incluida_no_painel_de_saude(s):
    assert relatorios.ultima_execucao(s) is None and relatorios.saude(s, agora=AGORA)["ultima_execucao"] is None
    inicio = datetime(2026, 10, 23, 12, 0, 0)
    s.add(Execucao(iniciada_em=inicio - timedelta(days=1), finalizada_em=inicio - timedelta(days=1), origem="manual", total=1))
    s.add(Execucao(iniciada_em=inicio, finalizada_em=inicio + timedelta(minutes=7), origem="assistida", total=12, ok=10, falhas=2,
                   links_corrigidos=1, em_revisao=1))
    s.add(Execucao(iniciada_em=inicio + timedelta(hours=1), origem="manual", total=3))  # ainda rodando
    s.commit()
    u = relatorios.saude(s, agora=AGORA)["ultima_execucao"]
    assert (u["origem"], u["total"], u["em_andamento"], u["duracao_s"]) == ("manual", 3, True, None)

    s.query(Execucao).filter(Execucao.finalizada_em.is_(None)).delete()
    s.commit()
    u = relatorios.ultima_execucao(s)
    assert (u["total"], u["ok"], u["falhas"], u["duracao_s"], u["links_corrigidos"], u["em_revisao"]) == (12, 10, 2, 420, 1, 1)
    assert u["em_andamento"] is False
