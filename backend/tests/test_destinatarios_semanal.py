"""Destinatários editáveis pela interface, assinaturas por usuário e resumo semanal em Excel."""
import io
from datetime import datetime, timedelta

import pytest
from conftest import URL_S25
from openpyxl import load_workbook
from sqlalchemy import select

from searching import operacao
from searching.alerts import destinatarios as dest
from searching.alerts import semanal
from searching.alerts import servico as alertas
from searching.alerts.email import ConfigSMTP
from searching.alerts.servico import Notificador
from searching.core.coleta import _config
from searching.db.models import Assinatura, Coleta, Configuracao, Produto, Usuario
from searching.db.session import criar_engine, iniciar_banco

SMTP = ConfigSMTP("smtp.teste", 587, "u@t", "s", "remetente@t", ("equipe@t.com",))


def _notificador(enviadas):
    return Notificador(carregar_config=lambda: SMTP, enviar=lambda msg, cfg=None: enviadas.append(msg))


@pytest.fixture
def s(tmp_path):
    with iniciar_banco(criar_engine(tmp_path / "d.db"))() as sessao:
        yield sessao


def _produto(s, nome="Samsung Galaxy S25 256GB", url=URL_S25):
    p = Produto(nome_comercial=nome, loja="Magazine Luiza", url_atual=url)
    s.add(p)
    s.flush()
    return p


def _coleta(s, p, pix, quando, status="OK"):
    c = Coleta(produto_id=p.id, coletado_em=quando, status=status, preco_pix_centavos=pix)
    s.add(c)
    s.commit()
    return c


def _usuario(s, email, ativo=True):
    u = Usuario(email=email, nome=email.split("@")[0], papel="leitor", senha_hash="x", ativo=ativo)
    s.add(u)
    s.flush()
    return u


# ---------- lista de destinatários ----------
def test_validar_lista_de_destinatarios():
    assert dest.validar_lista(" A@x.com; b@y.com,\nA@X.com ") == ["a@x.com", "b@y.com"]
    assert dest.validar_lista([]) == []
    for ruim in ("sem-arroba", ["ok@x.com", "x@"], "a@b"):
        with pytest.raises(ValueError, match="inválido"):
            dest.validar_lista(ruim)
    with pytest.raises(ValueError, match="No máximo"):
        dest.validar_lista([f"u{i}@x.com" for i in range(dest.MAX_DESTINATARIOS + 1)])


def test_lista_da_interface_manda_no_env_e_vazia_volta_ao_env(s):
    assert dest.efetivos(s, ("env@x.com",)) == ("env@x.com",)
    dest.salvar(s, "novo@x.com, outro@x.com")
    assert dest.efetivos(s, ("env@x.com",)) == ("novo@x.com", "outro@x.com")
    dest.salvar(s, [])
    assert dest.efetivos(s, ("env@x.com",)) == ("env@x.com",)


def test_assinantes_somam_sem_repetir_e_ignoram_inativos(s):
    p, outro = _produto(s), _produto(s, "Outro", URL_S25 + "2")
    a, b, c = _usuario(s, "a@x.com"), _usuario(s, "b@x.com"), _usuario(s, "c@x.com", ativo=False)
    for u in (a, b, c):
        s.add(Assinatura(usuario_id=u.id, produto_id=p.id))
    s.add(Assinatura(usuario_id=b.id, produto_id=outro.id))
    s.commit()
    assert dest.efetivos(s, ("A@x.com", "env@x.com"), p.id) == ("A@x.com", "env@x.com", "b@x.com")
    assert dest.efetivos(s, ("env@x.com",), outro.id) == ("env@x.com", "b@x.com")
    assert dest.efetivos(s, ("env@x.com",)) == ("env@x.com",)  # sem produto: só a lista base


def test_alerta_de_queda_vai_para_a_lista_e_para_os_assinantes_do_produto(s):
    p, outro = _produto(s), _produto(s, "Outro", URL_S25 + "2")
    a = _usuario(s, "assinante@x.com")
    s.add(Assinatura(usuario_id=a.id, produto_id=p.id))
    dest.salvar(s, "gestor@x.com")
    antes = datetime.now() - timedelta(days=1)
    for prod in (p, outro):
        _coleta(s, prod, 100000, antes)
    enviadas = []
    for prod in (p, outro):
        c = _coleta(s, prod, 80000, datetime.now())
        alerta = alertas.processar_coleta(s, prod, c, _config(s), _notificador(enviadas))
        assert alerta is not None and alerta.sucesso
    assert enviadas[0]["To"] == "gestor@x.com, assinante@x.com"  # produto assinado
    assert enviadas[1]["To"] == "gestor@x.com"                   # produto sem assinante


def test_email_de_teste_usa_a_lista_da_interface(s):
    from searching.alerts.servico import enviar_teste

    dest.salvar(s, "gestor@x.com")
    enviadas = []
    ok, msg = enviar_teste(_notificador(enviadas), s=s)
    assert ok and enviadas[0]["To"] == "gestor@x.com" and "gestor@x.com" in msg


# ---------- API: destinatários (modo local e hospedado) ----------
def test_api_destinatarios_no_modo_local(tmp_path, monkeypatch):
    import importlib
    import sys

    monkeypatch.setenv("SEARCHING_DB", str(tmp_path / "l.db"))
    monkeypatch.delenv("SEARCHING_MODO", raising=False)
    sys.modules.pop("server", None)
    server = importlib.import_module("server")
    try:
        c = server.app.test_client()
        assert c.get("/api/alertas").get_json()["destinatarios"] == {"configurados": []}
        r = c.put("/api/alertas/destinatarios", json={"destinatarios": "Gestor@X.com, outro@y.com"})
        assert r.status_code == 200 and r.get_json()["configurados"] == ["gestor@x.com", "outro@y.com"]
        assert c.get("/api/alertas").get_json()["destinatarios"]["configurados"] == ["gestor@x.com", "outro@y.com"]
        assert c.put("/api/alertas/destinatarios", json={"destinatarios": ["ruim"]}).status_code == 400
    finally:
        sys.modules.pop("server", None)


def test_so_admin_edita_destinatarios_no_hospedado(hosp):
    corpo = {"destinatarios": ["gestor@x.com"]}
    assert hosp.cliente("admin@teste.com").put("/api/alertas/destinatarios", corpo).status_code == 200
    for quem in ("coletor@teste.com", "leitor@teste.com"):
        assert hosp.cliente(quem).put("/api/alertas/destinatarios", corpo).status_code == 403
    assert hosp.cliente("leitor@teste.com").get("/api/alertas").get_json()["destinatarios"]["configurados"] == ["gestor@x.com"]


# ---------- API: assinaturas ----------
def _produtos(hosp, n=2):
    adm = hosp.cliente("admin@teste.com")
    ids = [adm.post("/api/links", {"nome": f"Samsung Galaxy S{25 + i} 256GB", "link": URL_S25 + str(i)}).get_json()["id"] for i in range(n)]
    return adm, ids


def test_coletor_assina_produtos_e_leitor_nao_escreve(hosp):
    adm, (p1, p2) = _produtos(hosp)
    col, leitor = hosp.cliente("coletor@teste.com"), hosp.cliente("leitor@teste.com")
    assert col.get("/api/minhas-assinaturas").get_json() == {"produtos": []}
    r = col.put("/api/minhas-assinaturas", {"produtos": [p1, p2, 9999, "x"]})
    assert r.status_code == 200 and r.get_json()["produtos"] == [p1, p2]  # id inexistente é ignorado
    assert col.put("/api/minhas-assinaturas", {"produtos": [p2]}).get_json()["produtos"] == [p2]
    assert leitor.get("/api/minhas-assinaturas").status_code == 200
    assert leitor.put("/api/minhas-assinaturas", {"produtos": [p1]}).status_code == 403  # leitor continua sem escrever
    # o admin assina por ele
    uid = hosp.uid("leitor@teste.com")
    assert adm.put(f"/api/usuarios/{uid}/assinaturas", {"produtos": [p1]}).status_code == 200
    assert leitor.get("/api/minhas-assinaturas").get_json() == {"produtos": [p1]}
    mapa = adm.get("/api/assinaturas").get_json()
    assert mapa[str(uid)] == [p1] and col.get("/api/assinaturas").status_code == 403


# ---------- resumo semanal ----------
def _cenario_semanal(s):
    p1, p2 = _produto(s), _produto(s, "Motorola Edge 50 256GB", URL_S25 + "2")
    agora = datetime(2026, 10, 26, 9, 0)
    for dias, pix1, pix2 in ((6, 100000, 50000), (3, 90000, 50000), (0, 80000, 52000)):
        _coleta(s, p1, pix1, agora - timedelta(days=dias, minutes=1))
        _coleta(s, p2, pix2, agora - timedelta(days=dias, minutes=2))
    _coleta(s, p2, 1, agora - timedelta(days=30))  # fora da janela
    u = _usuario(s, "so-p1@x.com")
    s.add(Assinatura(usuario_id=u.id, produto_id=p1.id))
    s.commit()
    return p1, p2, agora


def _planilha(msg):
    anexo = next(msg.iter_attachments())
    assert anexo.get_filename().endswith(".xlsx")
    return load_workbook(io.BytesIO(anexo.get_content()))["Histórico"]


def test_resumo_semanal_envia_excel_de_tudo_para_a_lista_e_so_o_assinado_para_o_usuario(s):
    p1, p2, agora = _cenario_semanal(s)
    enviadas = []
    r = semanal.enviar(s, _notificador(enviadas), agora=agora)
    assert sorted(r["enviados"]) == ["equipe@t.com", "so-p1@x.com"] and r["falhas"] == {}
    por_destino = {m["To"]: m for m in enviadas}
    tudo, so_p1 = _planilha(por_destino["equipe@t.com"]), _planilha(por_destino["so-p1@x.com"])
    assert tudo.max_row - 1 == 6 and so_p1.max_row - 1 == 3  # 30 dias atrás fica de fora
    assert {linha[1] for linha in so_p1.iter_rows(min_row=2, values_only=True)} == {"Samsung Galaxy S25 256GB"}
    corpo = por_destino["equipe@t.com"].get_body(("plain",)).get_content()
    assert "Samsung Galaxy S25 256GB" in corpo and "-20,0%" in corpo  # 1000 -> 800: queda de 20% na semana
    assert corpo.index("Samsung Galaxy S25") < corpo.index("Motorola")  # maiores quedas primeiro


def test_resumo_semanal_respeita_a_semana_forcar_e_simular(s):
    _, _, agora = _cenario_semanal(s)
    enviadas = []
    n = _notificador(enviadas)
    assert semanal.enviar(s, n, agora=agora, simular=True)["enviados"] and enviadas == []  # simulação não envia nem marca
    assert semanal.enviar(s, n, agora=agora)["enviados"]
    antes = len(enviadas)
    r = semanal.enviar(s, n, agora=agora + timedelta(days=2))
    assert r["pulado"] and len(enviadas) == antes
    assert semanal.enviar(s, n, agora=agora + timedelta(days=2), forcar=True)["enviados"]
    p1 = s.scalar(select(Produto))
    _coleta(s, p1, 70000, agora + timedelta(days=7))
    assert semanal.enviar(s, n, agora=agora + timedelta(days=8, hours=12))["enviados"]  # 6,5 dias depois do último envio


def test_resumo_semanal_sem_coletas_nao_envia_e_falha_de_um_nao_trava_os_outros(s):
    enviadas = []
    assert semanal.enviar(s, _notificador(enviadas))["enviados"] == [] and enviadas == []
    _, _, agora = _cenario_semanal(s)

    def enviar(msg, cfg=None):
        if msg["To"] == "equipe@t.com":
            raise RuntimeError("SMTP caiu")
        enviadas.append(msg)

    r = semanal.enviar(s, Notificador(carregar_config=lambda: SMTP, enviar=enviar), agora=agora)
    assert r["enviados"] == ["so-p1@x.com"] and "SMTP caiu" in r["falhas"]["equipe@t.com"]


def test_lista_da_interface_tambem_vale_no_resumo_semanal(s):
    _, _, agora = _cenario_semanal(s)
    dest.salvar(s, "gestor@x.com")
    enviadas = []
    r = semanal.enviar(s, _notificador(enviadas), agora=agora)
    assert sorted(r["enviados"]) == ["gestor@x.com", "so-p1@x.com"]


def test_cli_resumo_semanal_so_roda_se_estiver_ligado(tmp_path, monkeypatch):
    monkeypatch.setenv("SEARCHING_DB", str(tmp_path / "c.db"))
    iniciar_banco(criar_engine())
    saidas = []
    assert operacao.main(["resumo-semanal"], saida=saidas.append) == 0 and "desligado" in saidas[0]
    with iniciar_banco(criar_engine())() as s:
        s.merge(Configuracao(chave="resumo_semanal_email", valor="1"))
        s.commit()
        assert s.scalar(select(Configuracao.valor).where(Configuracao.chave == "resumo_semanal_email")) == "1"
    assert operacao.main(["resumo-semanal", "--simular"], saida=saidas.append) in (0, 1)  # sem SMTP no .env: não explode
    assert "Resumo semanal" in saidas[-1] or "Erro" in saidas[-1] or "SMTP" in saidas[-1]
