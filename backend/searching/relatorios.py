"""Consultas de relatório sobre o banco: histórico filtrado, galeria de prints, exportação para Excel e painel de saúde."""
from datetime import date, datetime, timedelta
from io import BytesIO
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .core.precos import formatar_brl
from .core.status import StatusColeta
from .db.models import Alerta, Coleta, Execucao, Produto, RevisaoLink

LIMITE_EXPORTACAO = 50_000


def _data(valor) -> date | None:
    if not valor:
        return None
    if isinstance(valor, date):
        return valor
    return datetime.strptime(str(valor)[:10], "%Y-%m-%d").date()


def print_url(caminho: str | None) -> str | None:
    """URL (servida por /api/prints/<pasta>/<arquivo>) do print de uma coleta."""
    if not caminho:
        return None
    partes = caminho.replace("\\", "/").split("/")
    if len(partes) < 2:
        return None
    return f"/api/prints/{quote(partes[-2])}/{quote(partes[-1])}"


def _consulta(de=None, ate=None, loja=None, produto_id=None, status=None):
    consulta = select(Coleta, Produto).join(Produto, Coleta.produto_id == Produto.id)
    if _data(de):
        consulta = consulta.where(Coleta.coletado_em >= datetime.combine(_data(de), datetime.min.time()))
    if _data(ate):
        consulta = consulta.where(Coleta.coletado_em < datetime.combine(_data(ate) + timedelta(days=1), datetime.min.time()))
    if loja:
        consulta = consulta.where(Produto.loja == loja)
    if produto_id:
        consulta = consulta.where(Produto.id == produto_id)
    if status:
        consulta = consulta.where(Coleta.status == status)
    return consulta


def historico(s: Session, **filtros) -> list[dict]:
    """Coletas na forma que as telas Relatórios e Curvas consomem (texto formatado + valores numéricos)."""
    linhas = s.execute(_consulta(**filtros).order_by(Coleta.coletado_em, Coleta.id)).all()
    return [
        {
            "ID": c.id,
            "EXECUCAO_ID": c.execucao_id,
            "DATA": c.coletado_em.strftime("%Y-%m-%d %H:%M:%S"),
            "NOME": p.nome_comercial,
            "MODELO": c.titulo_lido or p.nome_comercial,
            "SITE": p.loja,
            "ENTREGUE_POR": c.vendedor or "",
            "PIX": formatar_brl(c.preco_pix_centavos) if c.preco_pix_centavos is not None else "",
            "CREDITO": formatar_brl(c.preco_cartao_1x_centavos) if c.preco_cartao_1x_centavos is not None else "",
            "STATUS": c.status,
            "LINK": p.url_atual,
            "PRODUTO_ID": p.id,
            "PIX_CENTAVOS": c.preco_pix_centavos,
            "CREDITO_CENTAVOS": c.preco_cartao_1x_centavos,
            "ERRO_DETALHE": c.erro_detalhe or "",
            "PRINT": c.caminho_print or "",
        }
        for c, p in linhas
    ]


# ---------- galeria de prints ----------
def datas_com_prints(s: Session) -> list[dict]:
    dia = func.date(Coleta.coletado_em)
    linhas = s.execute(
        select(dia, func.count(Coleta.id), func.count(Coleta.caminho_print)).group_by(dia).order_by(dia.desc())
    ).all()
    return [{"data": d, "total": total, "com_print": com_print} for d, total, com_print in linhas if com_print]


def galeria(s: Session, data=None, produto_id=None) -> list[dict]:
    """Prints de uma data e/ou de um produto, cada um com o preço coletado naquele momento."""
    consulta = _consulta(de=data, ate=data, produto_id=produto_id).where(Coleta.caminho_print.is_not(None))
    linhas = s.execute(consulta.order_by(Coleta.coletado_em.desc(), Coleta.id.desc())).all()
    return [
        {
            "id": c.id,
            "produto_id": p.id,
            "produto": p.nome_comercial,
            "loja": p.loja,
            "coletado_em": c.coletado_em.isoformat(timespec="seconds"),
            "status": c.status,
            "pix": formatar_brl(c.preco_pix_centavos) if c.preco_pix_centavos is not None else None,
            "cartao_1x": formatar_brl(c.preco_cartao_1x_centavos) if c.preco_cartao_1x_centavos is not None else None,
            "erro_detalhe": c.erro_detalhe,
            "print_url": print_url(c.caminho_print),
        }
        for c, p in linhas
    ]


# ---------- Excel ----------
COLUNAS_EXCEL = [
    ("Data", 18), ("Produto", 34), ("Título lido", 46), ("Loja", 16), ("Vendedor", 18), ("Status", 22),
    ("PIX (R$)", 13), ("Cartão 1x (R$)", 15), ("Link", 50), ("Print (caminho)", 52), ("Detalhe", 36),
]


def exportar_excel(s: Session, ids: list[int]) -> bytes:
    """Planilha das coletas pedidas, com preços numéricos e o caminho do print."""
    if not ids:
        raise ValueError("Nenhuma coleta para exportar")
    if len(ids) > LIMITE_EXPORTACAO:
        raise ValueError(f"Máximo de {LIMITE_EXPORTACAO} linhas por exportação")
    linhas = s.execute(
        select(Coleta, Produto).join(Produto, Coleta.produto_id == Produto.id).where(Coleta.id.in_(ids)).order_by(Coleta.coletado_em, Coleta.id)
    ).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Histórico"
    ws.append([nome for nome, _ in COLUNAS_EXCEL])
    for c, p in linhas:
        ws.append([
            c.coletado_em, p.nome_comercial, c.titulo_lido or "", p.loja, c.vendedor or "", c.status,
            c.preco_pix_centavos / 100 if c.preco_pix_centavos is not None else None,
            c.preco_cartao_1x_centavos / 100 if c.preco_cartao_1x_centavos is not None else None,
            p.url_atual, c.caminho_print or "", c.erro_detalhe or "",
        ])

    cabecalho = PatternFill("solid", fgColor="1F3A5F")
    for i, (_, largura) in enumerate(COLUNAS_EXCEL, start=1):
        celula = ws.cell(row=1, column=i)
        celula.font, celula.fill = Font(bold=True, color="FFFFFF"), cabecalho
        celula.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = largura
    for linha in ws.iter_rows(min_row=2):
        linha[0].number_format = "dd/mm/yyyy hh:mm"
        linha[6].number_format = linha[7].number_format = '"R$" #,##0.00'
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------- painel de saúde ----------
def ultima_execucao(s: Session) -> dict | None:
    """Resumo da execução mais recente (automática ou assistida)."""
    e = s.scalars(select(Execucao).order_by(Execucao.iniciada_em.desc()).limit(1)).first()
    if e is None:
        return None
    duracao = (e.finalizada_em - e.iniciada_em).total_seconds() if e.finalizada_em else None
    return {
        "id": e.id,
        "iniciada_em": e.iniciada_em.isoformat(timespec="seconds"),
        "duracao_s": round(duracao) if duracao is not None else None,
        "origem": e.origem,
        "total": e.total,
        "ok": e.ok,
        "falhas": e.falhas,
        "links_corrigidos": e.links_corrigidos,
        "em_revisao": e.em_revisao,
        "em_andamento": e.finalizada_em is None,
    }


def saude(s: Session, agora: datetime | None = None, dias: int = 7) -> dict:
    """Taxa de sucesso por loja, falhas por motivo, links com problema e alertas dos últimos dias."""
    agora = agora or datetime.now()
    desde = agora - timedelta(days=dias)

    # Uma coleta "final" por produto em cada execução (a refeita após corrigir o link vale mais que a falha anterior).
    finais: dict[tuple, tuple[Coleta, Produto]] = {}
    for c, p in s.execute(_consulta().where(Coleta.coletado_em >= desde).order_by(Coleta.coletado_em, Coleta.id)).all():
        finais[(c.execucao_id, p.id)] = (c, p)

    por_loja: dict[str, dict] = {}
    por_status: dict[str, int] = {}
    for c, p in finais.values():
        loja = por_loja.setdefault(p.loja or "—", {"loja": p.loja or "—", "total": 0, "ok": 0})
        loja["total"] += 1
        if c.status == StatusColeta.OK.value:
            loja["ok"] += 1
        else:
            por_status[c.status] = por_status.get(c.status, 0) + 1
    lojas = sorted(
        ({**l, "taxa_sucesso": round(l["ok"] / l["total"] * 100, 1)} for l in por_loja.values()), key=lambda l: l["loja"]
    )

    # Estado atual de cada produto ativo: a última coleta dele.
    com_falha = []
    for p in s.scalars(select(Produto).where(Produto.ativo.is_(True)).order_by(Produto.id)):
        ultima = s.scalars(select(Coleta).where(Coleta.produto_id == p.id).order_by(Coleta.coletado_em.desc(), Coleta.id.desc()).limit(1)).first()
        if ultima and ultima.status != StatusColeta.OK.value:
            com_falha.append({"produto": p.nome_comercial, "status": ultima.status})

    alertas = s.execute(select(Alerta.sucesso, func.count()).where(Alerta.enviado_em >= desde).group_by(Alerta.sucesso)).all()
    contagem = {bool(ok): n for ok, n in alertas}
    return {
        "periodo_dias": dias,
        "lojas": lojas,
        "falhas_por_status": por_status,
        "produtos_com_falha_na_ultima_coleta": com_falha,
        "links_em_revisao": s.scalar(select(func.count()).select_from(RevisaoLink).where(RevisaoLink.estado == "pendente")) or 0,
        "alertas": {"enviados": contagem.get(True, 0), "falhos": contagem.get(False, 0)},
        "ultima_execucao": ultima_execucao(s),
    }
